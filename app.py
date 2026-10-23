import io
import os
import time
from datetime import datetime
from flask import Flask, g, jsonify, render_template, request, send_file

try:
    import psycopg2
    import psycopg2.extras
except ImportError:
    psycopg2 = None

import sqlite3

# Librería para generar hojas de Excel profesionales
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATABASE_URL = os.environ.get("DATABASE_URL")

app = Flask(__name__)

def get_db():
    if "db" not in g:
        if DATABASE_URL:
            url = DATABASE_URL.replace("postgres://", "postgresql://", 1) if DATABASE_URL.startswith("postgres://") else DATABASE_URL
            g.db = psycopg2.connect(url, cursor_factory=psycopg2.extras.DictCursor)
        else:
            DB_PATH = os.path.join(BASE_DIR, "inventario.db")
            g.db = sqlite3.connect(DB_PATH)
            g.db.row_factory = sqlite3.Row
    return g.db

@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()

def init_db():
    conn = get_db()
    cursor = conn.cursor()
    if DATABASE_URL:
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS tvs (
                id            SERIAL PRIMARY KEY,
                room          TEXT NOT NULL,
                ubicacion     TEXT,
                marca         TEXT,
                modelo        TEXT NOT NULL,
                serial        TEXT NOT NULL,
                perifericos   TEXT,
                instalacion   TEXT,
                conexion      TEXT,
                intergo       TEXT,
                comentarios   TEXT,
                createdat     BIGINT NOT NULL
            )
            """
        )
    else:
        DB_PATH = os.path.join(BASE_DIR, "inventario.db")
        cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS tvs (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                room          TEXT NOT NULL,
                ubicacion     TEXT,
                marca         TEXT,
                modelo        TEXT NOT NULL,
                serial        TEXT NOT NULL,
                perifericos   TEXT,
                instalacion   TEXT,
                conexion      TEXT,
                intergo       TEXT,
                comentarios   TEXT,
                createdat     INTEGER NOT NULL
            )
            """
        )
    
    if DATABASE_URL:
        cursor.execute("ALTER TABLE tvs ADD COLUMN IF NOT EXISTS conexion TEXT")
        cursor.execute("ALTER TABLE tvs ADD COLUMN IF NOT EXISTS intergo TEXT")
        cursor.execute("ALTER TABLE tvs ADD COLUMN IF NOT EXISTS comentarios TEXT")
    else:
        cursor.execute("PRAGMA table_info(tvs)")
        columns = [row[1] for row in cursor.fetchall()]
        if "conexion" not in columns:
            cursor.execute("ALTER TABLE tvs ADD COLUMN conexion TEXT")
        if "intergo" not in columns:
            cursor.execute("ALTER TABLE tvs ADD COLUMN intergo TEXT")
        if "comentarios" not in columns:
            cursor.execute("ALTER TABLE tvs ADD COLUMN comentarios TEXT")
            
    conn.commit()
    cursor.close()

with app.app_context():
    init_db()

FIELDS = [
    "room",
    "ubicacion",
    "marca",
    "modelo",
    "serial",
    "perifericos",
    "conexion",
    "intergo",
    "comentarios",
    "instalacion",
]

def format_datetime(ts):
    if not ts:
        return "—"
    try:
        dt = datetime.fromtimestamp(ts / 1000.0)
        return dt.strftime("%d/%m/%Y · %H:%M")
    except Exception:
        return "—"

def row_to_dict(row) -> dict:
    d = dict(row)
    d["id"] = str(d["id"])
    if "createdat" in d:
        d["createdAt"] = d["createdat"]
    return d

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/api/tvs", methods=["GET"])
def list_tvs():
    db = get_db()
    cursor = db.cursor()
    cursor.execute("SELECT * FROM tvs ORDER BY createdat DESC")
    rows = cursor.fetchall()
    cursor.close()
    return jsonify([row_to_dict(r) for r in rows])

@app.route("/api/tvs/export", methods=["GET"])
def export_excel():
    db = get_db()
    cursor = db.cursor()
    cursor.execute("SELECT * FROM tvs ORDER BY createdat ASC")
    rows = cursor.fetchall()
    cursor.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario TV"

    # Mostrar líneas de cuadrícula en Excel
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Num. Hab.", "Ubicación", "Marca", "Modelo", "Serial",
        "Tipo de conexión", "¿Reconoce periféricos?", "InterGO+",
        "Comentarios adicionales", "Tipo de instalación", "Fecha y hora"
    ]
    ws.append(headers)

    # Estilos visuales
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366690", end_color="366690", fill_type="solid")
    
    # Fondo blanco uniforme para todas las filas
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    # Border delgado y limpio para definir las celdas
    thin_border = Border(
        left=Side(style='thin', color='D7DEE7'),
        right=Side(style='thin', color='D7DEE7'),
        top=Side(style='thin', color='D7DEE7'),
        bottom=Side(style='thin', color='D7DEE7')
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # Aplicar estilos a encabezados
    ws.row_dimensions[1].height = 26
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    # Cargar los datos
    for row_idx, r in enumerate(rows, start=2):
        item = row_to_dict(r)
        row_data = [
            item.get("room", ""),
            item.get("ubicacion", ""),
            item.get("marca", ""),
            item.get("modelo", ""),
            item.get("serial", ""),
            item.get("conexion", ""),
            item.get("perifericos", ""),
            item.get("intergo", ""),
            item.get("comentarios", ""),
            item.get("instalacion", ""),
            format_datetime(item.get("createdAt"))
        ]
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 21

        # Formatear celdas de la fila (Fondo blanco + bordes delgados)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.fill = white_fill
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            
            # Centrar campos de selección y código, alinear texto largo a la izquierda
            if col_idx in [1, 2, 3, 6, 7, 8, 10, 11]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Ajuste automático del ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Crear respuesta para descarga directa en el navegador
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="Inventario_TVs.xlsx"
    )

@app.route("/api/tvs", methods=["POST"])
def create_tv():
    payload = request.get_json(silent=True) or {}

    room = (payload.get("room") or "").strip()
    modelo = (payload.get("modelo") or "").strip()
    serial = (payload.get("serial") or "").strip()
    if not room or not modelo or not serial:
        return jsonify({"error": "room, modelo y serial son obligatorios"}), 400

    values = {field: (payload.get(field) or "") for field in FIELDS}
    values["room"] = room
    values["modelo"] = modelo
    values["serial"] = serial
    created_at = int(time.time() * 1000)

    columns = FIELDS + ["createdat"]
    columns_str = ", ".join(columns)
    
    ph = "%s" if DATABASE_URL else "?"
    placeholders = ", ".join([ph] * len(columns))

    db = get_db()
    cursor = db.cursor()

    if DATABASE_URL:
        cursor.execute(
            f"INSERT INTO tvs ({columns_str}) VALUES ({placeholders}) RETURNING id",
            [values[f] for f in FIELDS] + [created_at],
        )
        new_id = cursor.fetchone()["id"]
        db.commit()
        cursor.execute("SELECT * FROM tvs WHERE id = %s", (new_id,))
        new_row = cursor.fetchone()
    else:
        cursor.execute(
            f"INSERT INTO tvs ({columns_str}) VALUES ({placeholders})",
            [values[f] for f in FIELDS] + [created_at],
        )
        db.commit()
        cursor.execute("SELECT * FROM tvs WHERE id = ?", (cursor.lastrowid,))
        new_row = cursor.fetchone()

    cursor.close()
    return jsonify(row_to_dict(new_row)), 201

@app.route("/api/tvs/<int:tv_id>", methods=["PUT"])
def update_tv(tv_id):
    payload = request.get_json(silent=True) or {}

    db = get_db()
    cursor = db.cursor()
    
    sel_ph = "%s" if DATABASE_URL else "?"
    cursor.execute(f"SELECT * FROM tvs WHERE id = {sel_ph}", (tv_id,))
    existing = cursor.fetchone()
    if existing is None:
        cursor.close()
        return jsonify({"error": "No encontrado"}), 404

    room = (payload.get("room") or "").strip()
    modelo = (payload.get("modelo") or "").strip()
    serial = (payload.get("serial") or "").strip()
    if not room or not modelo or not serial:
        cursor.close()
        return jsonify({"error": "room, modelo y serial son obligatorios"}), 400

    values = {field: (payload.get(field) or "") for field in FIELDS}
    values["room"] = room
    values["modelo"] = modelo
    values["serial"] = serial

    eq_ph = "%s" if DATABASE_URL else "?"
    set_clause = ", ".join(f"{f} = {eq_ph}" for f in FIELDS)
    
    cursor.execute(
        f"UPDATE tvs SET {set_clause} WHERE id = {eq_ph}",
        [values[f] for f in FIELDS] + [tv_id],
    )
    db.commit()
    
    cursor.execute(f"SELECT * FROM tvs WHERE id = {sel_ph}", (tv_id,))
    updated = cursor.fetchone()
    cursor.close()
    return jsonify(row_to_dict(updated))

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)